"""The workbook `scrape --export-dir` writes: the export's own rows, marked by the run's own diff.

`scrapewatch.cli._export_workbook` is the glue between a run and the workbook
module. It reads each collected source's latest snapshot the way the exports do,
takes its marks from the diff the run made, and leaves out whatever the run did
not collect. These tests drive it against a real SQLite file, so the snapshots,
the diff and the export the sheet is held to are the pipeline's own.
"""
from __future__ import annotations

import json
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from scrapewatch.cli import SHEET_TITLES, WORKBOOK_FILE, _export_workbook
from scrapewatch.models import Record
from scrapewatch.pipeline.diff import diff
from scrapewatch.pipeline.report import ChangeReport
from scrapewatch.storage import Storage
from scrapewatch.workbook import CHANGES_TITLE, COLLECTED_AT

pytestmark = pytest.mark.integration

COLLECTED = datetime(2026, 9, 22, 5, 12, tzinfo=UTC)
DEMO = SHEET_TITLES["demo"]


def _product(eid: str, name: str, price: str, stock: int) -> Record:
    return Record(
        source="demo", kind="product", external_id=eid, fetched_at=COLLECTED, url="http://127.0.0.1:8765/?page=1",
        fields={"name": name, "price": Decimal(price), "currency": "USD", "in_stock": stock > 0, "stock": stock},
    )


def _book(eid: str, title: str) -> Record:
    return Record(
        source="books", kind="book", external_id=eid, fetched_at=COLLECTED, url=f"https://b/{eid}",
        fields={"title": title, "price": Decimal("51.77"), "currency": "GBP", "in_stock": True, "rating": 3},
    )


@pytest.fixture
def storage(tmp_path: Path) -> Storage:
    s = Storage(f"sqlite:///{tmp_path / 'db.sqlite3'}")
    s.open()
    return s


def _night(storage: Storage, records: dict[str, list[Record]]) -> ChangeReport:
    """One run: a snapshot per source, each diffed against the one before it, as `run_sources` does."""
    run = storage.start_run(list(records))
    changesets = {}
    for source, snapshot in records.items():
        storage.save_snapshot(run, source, snapshot)
        latest, *previous = storage.latest_snapshots(source, n=2)
        changesets[source] = diff(previous[0] if previous else [], latest)
    storage.finish_run(run, {})
    return ChangeReport.from_changesets(changesets, generated_at=COLLECTED)


def _collected(*sources: str) -> dict[str, dict]:
    return {source: {"skipped": False} for source in sources}


def _filled(sheet) -> set[str]:
    return {cell.coordinate for row in sheet.iter_rows() for cell in row if cell.fill.fill_type == "solid"}


def test_the_sheet_says_what_the_json_export_of_the_same_snapshot_says(storage: Storage, tmp_path: Path) -> None:
    """Same columns, same order, same values, row for row: the workbook is written
    from the export's own layout, so it cannot disagree with the files beside it."""
    report = _night(storage, {"demo": [_product("2", "Basalt Coffee Mug", "8.99", 27),
                                       _product("1", "Aurora Desk Lamp", "24.99", 0)]})
    workbook = load_workbook(_export_workbook(storage, _collected("demo"), report, tmp_path))
    storage.export("demo", "json", tmp_path / "demo.json")
    exported = json.loads((tmp_path / "demo.json").read_text(encoding="utf-8"))

    sheet = workbook[DEMO]
    assert [cell.value for cell in sheet[1]] == [*exported[0], COLLECTED_AT]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(exported)
    for row, record in zip(rows, exported, strict=True):
        for value, (column, expected) in zip(row, record.items(), strict=False):
            if column == "price":
                assert Decimal(str(value)) == Decimal(expected), "a price is the same amount, as a number"
            else:
                assert value == expected, column
        assert row[-1] == datetime(2026, 9, 22, 5, 12)


def test_the_marks_are_the_diff_the_run_made(storage: Storage, tmp_path: Path) -> None:
    _night(storage, {"demo": [_product("1", "Aurora Desk Lamp", "24.99", 14),
                              _product("3", "Cobalt Backpack", "54.99", 9)]})
    report = _night(storage, {"demo": [_product("1", "Aurora Desk Lamp", "22.49", 14),
                                       _product("4", "Driftwood Cutting Board", "18.99", 0)]})
    workbook = load_workbook(_export_workbook(storage, _collected("demo"), report, tmp_path))

    sheet = workbook[DEMO]
    header = [cell.value for cell in sheet[1]]
    price = get_column_letter(header.index("price") + 1)
    new_row = {f"{get_column_letter(column)}3" for column in range(1, len(header) + 1)}
    assert _filled(sheet) == {f"{price}2"} | new_row, "the changed price, and the row of the new product"

    changes = [row[:6] for row in workbook[CHANGES_TITLE].iter_rows(min_row=2, values_only=True)]
    assert [row[:4] for row in changes] == [
        (DEMO, "1", "changed", "price"),
        (DEMO, "4", "added", None),
        (DEMO, "3", "removed", None),
    ]
    assert changes[0][4:] == (pytest.approx(24.99), pytest.approx(22.49))
    assert "name: Cobalt Backpack" in changes[2][4], "a product that is gone is still named"


def test_a_source_with_no_earlier_snapshot_is_marked_with_nothing(storage: Storage, tmp_path: Path) -> None:
    """The run's diff of a first snapshot calls every record new; the workbook does
    not repeat that as news, and says why instead."""
    report = _night(storage, {"demo": [_product("1", "Aurora Desk Lamp", "24.99", 14)]})
    assert report.changesets["demo"].added, "precondition: the diff called the product new"

    workbook = load_workbook(_export_workbook(storage, _collected("demo"), report, tmp_path))
    assert _filled(workbook[DEMO]) == set()
    changes = workbook[CHANGES_TITLE]
    assert changes["A2"].value.startswith(f"First snapshot of {DEMO}:")
    assert changes.max_row == 2


def test_a_source_the_run_did_not_collect_gets_no_sheet(storage: Storage, tmp_path: Path) -> None:
    """The sheets follow the run's own verdict and its order, the way the exports do."""
    report = _night(storage, {"books": [_book("b1", "Sharp Objects")], "demo": [_product("1", "Lamp", "24.99", 1)]})
    stats = {"books": {"skipped": False}, "quotes": {"skipped": True}, "demo": {"skipped": False}}
    workbook = load_workbook(_export_workbook(storage, stats, report, tmp_path))
    assert workbook.sheetnames == [SHEET_TITLES["books"], DEMO, CHANGES_TITLE]


def test_a_run_that_collected_nothing_writes_no_workbook(storage: Storage, tmp_path: Path) -> None:
    empty = ChangeReport.from_changesets({}, generated_at=COLLECTED)
    assert _export_workbook(storage, {"demo": {"skipped": True}}, empty, tmp_path) is None
    assert not (tmp_path / WORKBOOK_FILE).exists()
