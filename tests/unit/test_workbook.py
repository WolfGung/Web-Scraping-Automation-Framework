"""The workbook a client opens: the right sheets, real numbers, and marks only where something moved.

`scrapewatch.workbook.write_workbook` is handed plain data here — two sources'
tables and the night's changes, one changed price, one book that is new and one
that is gone — and the file it writes is opened again with `openpyxl`, the way a
reader's spreadsheet would open it. Everything asserted is a property of that
file: sheet names and order, header cells, row counts, the frozen header and the
filter, column widths, number formats, value types, links, which cells carry a
fill, the file's own properties, and what each sheet tells Excel to ignore.
"""
from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from scrapewatch.cli import SHEET_TITLES
from scrapewatch.sources import KNOWN_SOURCES
from scrapewatch.workbook import (
    CHANGES_HEADER,
    CHANGES_TITLE,
    COLLECTED_AT,
    MAX_WIDTH,
    NO_CHANGES,
    TOKEN_WIDTH,
    Change,
    Sheet,
    write_workbook,
)

pytestmark = pytest.mark.unit

COLLECTED = datetime(2026, 9, 22, 5, 12, 30, tzinfo=UTC)

ATTIC = "a-light-in-the-attic_1000"
VELVET = "tipping-the-velvet_999"
ATTIC_URL = f"https://books.toscrape.com/catalogue/{ATTIC}/index.html"
DEMO_URL = "http://127.0.0.1:8765/?page=1"

BOOKS = Sheet(
    title="Books",
    header=["currency", "in_stock", "price", "rating", "title", "external_id", "url", COLLECTED_AT],
    rows=[
        ["GBP", True, Decimal("51.77"), 3, "A Light in the Attic", ATTIC, ATTIC_URL, COLLECTED],
        ["GBP", False, Decimal("53.74"), 1, "Tipping the Velvet", VELVET, f"https://b/{VELVET}", COLLECTED],
    ],
    changed=frozenset({(ATTIC, "price")}),
    added=frozenset({VELVET}),
)
DEMO = Sheet(
    title="Demo store",
    header=["currency", "in_stock", "name", "price", "stock", "external_id", "url", COLLECTED_AT],
    rows=[["USD", True, "Aurora Desk Lamp", Decimal("24.99"), 14, "1", DEMO_URL, COLLECTED]],
)
CHANGES = [
    Change("Books", ATTIC, "changed", "price", Decimal("49.99"), Decimal("51.77"), currency="GBP"),
    Change("Books", VELVET, "added", after={"title": "Tipping the Velvet", "price": Decimal("53.74")}, currency="GBP"),
    Change("Books", "soumission_998", "removed", before={"title": "Soumission", "price": Decimal("50.10")},
           currency="GBP"),
]

#: Where the `Changes` table starts: the legend has the top row to itself.
CHANGES_TOP = 2


def _written(tmp_path: Path, sheets=(BOOKS, DEMO), changes=CHANGES, **options):
    path = tmp_path / "scrapewatch.xlsx"
    write_workbook(path, list(sheets), list(changes), **options)
    return load_workbook(path)


@pytest.fixture
def workbook(tmp_path: Path):
    return _written(tmp_path)


def _top(sheet) -> int:
    return CHANGES_TOP if sheet.title == CHANGES_TITLE else 1


def _cell(sheet, row: int, column_name: str):
    """The cell in `row` under the header `column_name`."""
    header = [cell.value for cell in sheet[_top(sheet)]]
    return sheet.cell(row=row, column=header.index(column_name) + 1)


def _filled(sheet) -> set[str]:
    """Every cell in the sheet that carries a solid fill."""
    return {
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.fill is not None and cell.fill.fill_type == "solid"
    }


def _widths(sheet) -> dict[int, float | None]:
    """The width stored for every column up to the last one in use, by column index."""
    stored: dict[int, float | None] = {}
    for dimension in sheet.column_dimensions.values():
        for index in range(dimension.min, dimension.max + 1):
            stored[index] = dimension.width if dimension.customWidth else None
    return {index: stored.get(index) for index in range(1, sheet.max_column + 1)}


# -- the shape of the file --------------------------------------------------------


def test_the_sheets_are_the_sources_in_run_order_then_the_changes(workbook) -> None:
    assert workbook.sheetnames == ["Books", "Demo store", CHANGES_TITLE]


def test_each_sheet_starts_with_its_header_in_bold(workbook) -> None:
    for sheet, header in ((BOOKS.title, BOOKS.header), (DEMO.title, DEMO.header), (CHANGES_TITLE, CHANGES_HEADER)):
        first_row = workbook[sheet][_top(workbook[sheet])][: len(header)]
        assert [cell.value for cell in first_row] == list(header), sheet
        assert all(cell.font.bold for cell in first_row), f"{sheet}: the header is not bold"


def test_each_sheet_holds_one_row_per_record_or_change(workbook) -> None:
    assert workbook["Books"].max_row - 1 == len(BOOKS.rows)
    assert workbook["Demo store"].max_row - 1 == len(DEMO.rows)
    assert workbook[CHANGES_TITLE].max_row - CHANGES_TOP == len(CHANGES)


def test_the_header_is_frozen_and_the_filter_covers_every_row(workbook) -> None:
    expected = {
        "Books": ("A2", f"A1:{get_column_letter(len(BOOKS.header))}{1 + len(BOOKS.rows)}"),
        "Demo store": ("A2", f"A1:{get_column_letter(len(DEMO.header))}{1 + len(DEMO.rows)}"),
        CHANGES_TITLE: ("A3", f"A2:{get_column_letter(len(CHANGES_HEADER))}{CHANGES_TOP + len(CHANGES)}"),
    }
    for name, (frozen, ref) in expected.items():
        assert workbook[name].freeze_panes == frozen, name
        assert workbook[name].auto_filter.ref == ref, name


def test_every_column_has_a_width_and_none_is_wider_than_the_cap(workbook) -> None:
    """A width is the longest value the column shows, header included, and a long
    title stops at the cap instead of turning the sheet into a wall."""
    for sheet in workbook.worksheets:
        widths = _widths(sheet)
        assert all(width is not None for width in widths.values()), f"{sheet.title}: {widths}"
        assert all(width <= MAX_WIDTH for width in widths.values()), f"{sheet.title}: {widths}"
    books = _widths(workbook["Books"])
    header = list(BOOKS.header)
    assert books[header.index("rating") + 1] < books[header.index("title") + 1], "widths follow the content"


def test_a_column_of_addresses_or_slugs_stops_at_the_narrower_cap(tmp_path: Path) -> None:
    """Text with no space to break at never wraps, so its column is clipped
    rather than widened: slugs and addresses stop at the narrower cap, and the
    last columns stay on the screen. Prose keeps the wider one and wraps."""
    prose = "It is our choices, Harry, that show what we truly are, far more than our abilities."
    sheet = Sheet(
        title="Quotes",
        header=["text", "external_id", "url"],
        rows=[[prose, "a-slug-" * 12, ATTIC_URL * 2]],
    )
    widths = _widths(_written(tmp_path, sheets=[sheet], changes=[])["Quotes"])
    assert widths[1] == MAX_WIDTH, "prose"
    assert widths[2] == TOKEN_WIDTH, "a slug"
    assert widths[3] == TOKEN_WIDTH, "an address"


# -- the values, as a spreadsheet reads them ---------------------------------------


def test_a_price_is_a_number_with_two_decimals_and_its_currency_symbol(workbook) -> None:
    books_price = _cell(workbook["Books"], 2, "price")
    assert isinstance(books_price.value, int | float) and not isinstance(books_price.value, bool)
    assert books_price.value == pytest.approx(51.77)
    assert books_price.number_format == '"£"#,##0.00'
    assert _cell(workbook["Demo store"], 2, "price").number_format == '"$"#,##0.00'


def test_a_price_on_the_changes_sheet_carries_its_currency_too(workbook) -> None:
    changes = workbook[CHANGES_TITLE]
    for column in ("before", "after"):
        price = _cell(changes, CHANGES_TOP + 1, column)
        assert isinstance(price.value, int | float), column
        assert price.number_format == '"£"#,##0.00', column


def test_a_count_is_an_integer_a_flag_a_boolean_and_a_time_a_date(workbook) -> None:
    stock = _cell(workbook["Demo store"], 2, "stock")
    assert stock.value == 14 and stock.number_format == "0"
    rating = _cell(workbook["Books"], 2, "rating")
    assert rating.value == 3 and rating.number_format == "0"
    assert _cell(workbook["Books"], 2, "in_stock").value is True
    assert _cell(workbook["Books"], 3, "in_stock").value is False
    collected = _cell(workbook["Books"], 2, COLLECTED_AT)
    assert collected.value == datetime(2026, 9, 22, 5, 12, 30), "naive UTC, as the header says"
    assert collected.number_format == "yyyy-mm-dd hh:mm"


def test_a_public_address_is_a_link_and_a_loopback_one_is_not(workbook) -> None:
    """A book's page is somewhere a reader can go; the demo store's loopback
    address only ever answered on the machine that scraped it."""
    book = _cell(workbook["Books"], 2, "url")
    assert book.hyperlink is not None and book.hyperlink.target == ATTIC_URL
    assert book.value == ATTIC_URL
    assert _cell(workbook["Demo store"], 2, "url").hyperlink is None


def test_text_that_starts_like_a_formula_stays_text(tmp_path: Path) -> None:
    """A scraped title is somebody else's text. Written as it stands, a title
    beginning with `=` would be a formula that runs in the reader's spreadsheet —
    the same hole the CSV export closes with a leading quote."""
    sheet = Sheet(title="Books", header=["title", "external_id"], rows=[["=HYPERLINK(\"x\")", "b1"]])
    cell = _written(tmp_path, sheets=[sheet], changes=[])["Books"]["A2"]
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("x")'


def test_long_prose_wraps_and_a_long_address_does_not(tmp_path: Path) -> None:
    """A sentence longer than the cap breaks across lines, where a reader can follow
    it; an address has no space to break at, and wrapping it would only chop it."""
    prose = "It is our choices, Harry, that show what we truly are, far more than our abilities. " * 2
    sheet = Sheet(title="Quotes", header=["text", "external_id", "url"], rows=[[prose.strip(), "q1", ATTIC_URL * 2]])
    quotes = _written(tmp_path, sheets=[sheet], changes=[])["Quotes"]
    assert quotes["A2"].alignment.wrap_text is True
    assert not quotes["C2"].alignment.wrap_text


def test_a_character_a_workbook_cannot_hold_is_dropped(tmp_path: Path) -> None:
    """A control character in scraped text is legal in JSON and illegal in the XML
    a workbook is made of; it is dropped rather than failing the night's export."""
    sheet = Sheet(title="Books", header=["title", "external_id"], rows=[["Bell\x07 Jar", "b1"]])
    assert _written(tmp_path, sheets=[sheet], changes=[])["Books"]["A2"].value == "Bell Jar"


# -- what moved, and where it shows ------------------------------------------------


def test_only_the_changed_cell_and_the_new_row_are_filled(workbook) -> None:
    books = workbook["Books"]
    changed = _cell(books, 2, "price").coordinate
    new_row = {f"{get_column_letter(column)}3" for column in range(1, len(BOOKS.header) + 1)}
    assert _filled(books) == {changed} | new_row
    assert books[changed].fill.fgColor.rgb != books["A3"].fill.fgColor.rgb, "two marks, two colours"
    assert _filled(workbook["Demo store"]) == set(), "nothing moved in the demo store"


def test_the_changes_sheet_lists_each_change_and_colours_its_kind(workbook) -> None:
    changes = workbook[CHANGES_TITLE]
    first = CHANGES_TOP + 1
    rows = [[cell.value for cell in row[: len(CHANGES_HEADER)]] for row in changes.iter_rows(min_row=first)]
    assert rows[0] == ["Books", ATTIC, "changed", "price", pytest.approx(49.99), pytest.approx(51.77)]
    assert rows[1][:4] == ["Books", VELVET, "added", None] and rows[1][4] is None
    assert rows[1][5] == "title: Tipping the Velvet; price: £53.74", "an added record is described in words"
    assert rows[2][:4] == ["Books", "soumission_998", "removed", None] and rows[2][5] is None
    assert rows[2][4] == "title: Soumission; price: £50.10"
    kinds = {changes.cell(row=row, column=3).fill.fgColor.rgb for row in range(first, first + len(CHANGES))}
    assert len(kinds) == 3, "changed, added and removed each have a colour of their own"


def test_the_changes_sheet_says_once_above_the_table_what_the_fills_mean(workbook) -> None:
    changes = workbook[CHANGES_TITLE]
    legend = changes["A1"]
    assert legend.font.italic and not legend.alignment.wrap_text
    assert legend.value.startswith("Yellow cell: value changed since the previous snapshot. Green row: new record.")
    assert [cell.value for cell in changes[1][1:]] == [None] * (changes.max_column - 1), "one note, in A1"
    assert changes.max_column == len(CHANGES_HEADER), "no spacer columns beside the table"
    for data_sheet in (BOOKS, DEMO):
        assert workbook[data_sheet.title].max_column == len(data_sheet.header), (
            f"{data_sheet.title} carries no legend of its own"
        )


def test_a_first_night_says_so_in_a3_and_marks_nothing(tmp_path: Path) -> None:
    """With no earlier snapshot, a diff calls every record new. That is not news,
    so nothing is marked and the Changes sheet says why in one line."""
    plain = Sheet(title="Books", header=BOOKS.header, rows=BOOKS.rows)
    workbook = _written(tmp_path, sheets=[plain], changes=[], uncompared=["Books"])
    changes = workbook[CHANGES_TITLE]
    assert changes["A3"].value.startswith("First snapshot of Books:")
    assert changes.max_row == 3 and changes.auto_filter.ref == "A2:F2"
    assert _filled(workbook["Books"]) == set()


def test_a_night_without_changes_says_so_rather_than_leaving_the_table_empty(tmp_path: Path) -> None:
    """The change report states zero changes rather than omitting them, because
    silence reads as "not checked"; the workbook says it the same way."""
    changes = _written(tmp_path, sheets=[DEMO], changes=[])[CHANGES_TITLE]
    assert [cell.value for cell in changes[CHANGES_TOP][: len(CHANGES_HEADER)]] == list(CHANGES_HEADER)
    assert changes["A3"].value == NO_CHANGES
    assert changes.max_row == 3


def test_a_mixed_night_leaves_a_blank_row_before_the_first_snapshot_note(tmp_path: Path) -> None:
    """Books moved against yesterday; the demo store was collected for the first
    time. Its note sits under the changes, a row apart, not inside the table."""
    changes = _written(tmp_path, changes=CHANGES[:1], uncompared=["Demo store"])[CHANGES_TITLE]
    last_change = CHANGES_TOP + 1
    assert changes.cell(row=last_change, column=3).value == "changed"
    assert all(cell.value is None for cell in changes[last_change + 1])
    assert changes.cell(row=last_change + 2, column=1).value.startswith("First snapshot of Demo store:")
    assert changes.auto_filter.ref == f"A2:F{last_change}", "the note is not part of the filtered table"


# -- the file itself ---------------------------------------------------------------


def test_the_file_names_scrapewatch_as_its_author(workbook) -> None:
    """File, Info reads the creator and the title: this project, not the library it uses."""
    assert workbook.properties.creator == "ScrapeWatch"
    assert workbook.properties.title == "ScrapeWatch nightly data"


#: The children of a worksheet part, in the order the Office Open XML schema
#: (CT_Worksheet) requires them. A part is valid only if its children follow it.
WORKSHEET_ORDER = (
    "sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr",
    "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
    "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations",
    "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
    "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
    "legacyDrawingHF", "drawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts",
    "extLst",
)


def test_every_sheet_tells_excel_the_digits_in_its_ids_are_text_on_purpose(tmp_path: Path) -> None:
    """An id like `17` is text in the exports and stays text here, and Excel marks
    every such cell with a green triangle unless the sheet says to ignore it. The
    library cannot write that instruction, so the writer adds it to each sheet part,
    where the schema puts it, and the file must still open."""
    path = tmp_path / "scrapewatch.xlsx"
    write_workbook(path, [BOOKS, DEMO], CHANGES)
    with zipfile.ZipFile(path) as package:
        parts = sorted(name for name in package.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
        assert len(parts) == 3
        for part in parts:
            children = [element.tag.split("}")[-1] for element in ET.fromstring(package.read(part))]
            assert children.count("ignoredErrors") == 1, part
            positions = [WORKSHEET_ORDER.index(name) for name in children]
            assert positions == sorted(positions), f"{part}: {children} is not in schema order"
            ignored = ET.fromstring(package.read(part)).find("{*}ignoredErrors/{*}ignoredError")
            assert ignored.get("sqref") == "A1:XFD1048576" and ignored.get("numberStoredAsText") == "1"
    assert load_workbook(path).sheetnames == ["Books", "Demo store", CHANGES_TITLE]


def test_every_source_the_project_scrapes_has_a_sheet_title() -> None:
    """A source with no title of its own would reach the workbook under its code
    name, two with one title would collide, and `Changes` is taken."""
    assert set(SHEET_TITLES) == set(KNOWN_SOURCES)
    assert len(set(SHEET_TITLES.values())) == len(SHEET_TITLES)
    assert CHANGES_TITLE not in SHEET_TITLES.values()
