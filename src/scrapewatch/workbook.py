"""The night's data as one Excel workbook: a sheet per source, then what changed.

The CSV and JSON exports are for programs. The workbook is for a person who opens
it in a spreadsheet and wants to read, filter and sort the collection without
setting anything up, so it is the same data, formatted, rather than a new view of
it: each source's sheet carries the export's columns, in the export's order, with
the export's values (`scrapewatch.storage.export_layout` lays out both), and one
more column saying when each record was collected. Numbers are numbers: a price
has two decimals and the currency symbol it was scraped with, a count is an
integer, a time is a date, a flag is TRUE or FALSE, and a public page's address is
a link.

What moved since the previous snapshot is marked where it happened. A value that
changed gets a light yellow fill in its cell, and a record that is new tonight a
light green fill across its row. The `Changes` sheet lists every change, one row
per changed field or per added or removed record, and says once, above its table,
what the fills mean, so the data sheets stay plain tables that filter and sort
cleanly.

This module takes plain data and knows nothing about sources, storage or runs:
`scrapewatch.cli` gathers the night's tables and changes and hands them over.
"""
from __future__ import annotations

import io
import ipaddress
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scrapewatch.pipeline.normalize import CURRENCY_SYMBOLS

#: The column every data sheet ends with: when the record was collected, in UTC,
#: as a date a spreadsheet can sort and filter rather than as text.
COLLECTED_AT = "collected at (UTC)"

CHANGES_TITLE = "Changes"
CHANGES_HEADER = ("source", "record key", "change", "field", "before", "after")

#: What File, Info shows for the file: this project, not the library that wrote it.
CREATOR = "ScrapeWatch"
TITLE = "ScrapeWatch nightly data"

#: The widest a column of prose may grow, in characters; prose longer than that
#: wraps. Text with no space to break at — an address, a slug — never wraps, so
#: its column stops at the narrower cap and the value is clipped at the edge,
#: rather than pushing the last columns off the screen.
MAX_WIDTH = 60
TOKEN_WIDTH = 40

MONEY_FORMAT = "#,##0.00"
INTEGER_FORMAT = "0"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"

#: The two marks the data sheets carry, and the colour the `Changes` sheet gives
#: each kind of change — the same yellow and green, so one legend covers both.
CHANGED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
ADDED_FILL = PatternFill(fill_type="solid", fgColor="E2EFDA")
REMOVED_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
KIND_FILLS = {"changed": CHANGED_FILL, "added": ADDED_FILL, "removed": REMOVED_FILL}

LEGEND = (
    "Yellow cell: value changed since the previous snapshot. Green row: new record. "
    "Red, in the change column: record removed."
)
NO_CHANGES = "No changes since the previous snapshot."

#: The column whose values are the addresses of the pages the records came from.
URL_COLUMN = "url"

#: Room for the filter button a header cell carries, and for the cell's margins.
_FILTER_BUTTON = 2
_PADDING = 2

_HEADER_FONT = Font(bold=True)
_NOTE_FONT = Font(italic=True, color="595959")
_LINK_FONT = Font(color="0563C1", underline="single")
_TOP = Alignment(vertical="top")
_TOP_WRAPPED = Alignment(vertical="top", wrap_text=True)

_SYMBOLS = {code: symbol for symbol, code in CURRENCY_SYMBOLS.items()}

#: Excel marks every cell holding digits as text with a green triangle, and an id
#: like `17` is text here on purpose, as it is in the exports. This is the
#: instruction Excel itself writes when a reader clicks "Ignore error", for the
#: whole sheet. openpyxl cannot write it, so it is added to each sheet part after
#: the file is saved.
_IGNORED_ERRORS = b'<ignoredErrors><ignoredError sqref="A1:XFD1048576" numberStoredAsText="1"/></ignoredErrors>'
_WORKSHEET_PART = re.compile(r"xl/worksheets/sheet\d+\.xml")
#: What the schema (CT_Worksheet) lets follow `ignoredErrors` in a sheet: it goes
#: in front of the first of these, or at the end when there is none.
_FOLLOWS_IGNORED_ERRORS = re.compile(
    rb"<(?:smartTags|drawing|legacyDrawing|legacyDrawingHF|drawingHF|picture|oleObjects|controls"
    rb"|webPublishItems|tableParts|extLst)[\s/>]"
)


@dataclass(frozen=True)
class Sheet:
    """One source's data, laid out as its exports are, and what moved in it tonight.

    `changed` holds a `(key, column)` pair for every value that changed, and
    `added` the key of every record that is new; `key` names the column that holds
    a row's key.
    """

    title: str
    header: Sequence[str]
    rows: Sequence[Sequence[Any]]
    key: str = "external_id"
    changed: frozenset[tuple[str, str]] = frozenset()
    added: frozenset[str] = frozenset()


@dataclass(frozen=True)
class Change:
    """One row of the `Changes` sheet: a changed field, or a record added or removed.

    An added record's fields go in `after` and a removed one's in `before`, and the
    sheet writes them out in words, so a record that is gone can still be read.
    `currency` is the record's own, so a price shows the symbol it was scraped with.
    """

    source: str
    key: str
    change: Literal["changed", "added", "removed"]
    field: str | None = None
    before: Any = None
    after: Any = None
    currency: str | None = None


def write_workbook(
    path: str | Path,
    sheets: Sequence[Sheet],
    changes: Sequence[Change],
    *,
    uncompared: Sequence[str] = (),
) -> None:
    """Write `sheets` in the order given, then the `Changes` sheet, to `path`.

    `uncompared` names the sheets whose source had no earlier snapshot. A diff with
    nothing on one side calls every record new, which is not news, so the caller
    marks nothing in those sheets and the `Changes` sheet says why in a note.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = CREATOR
    workbook.properties.title = TITLE
    for sheet in sheets:
        _write_data_sheet(workbook.create_sheet(sheet.title), sheet)
    compared = [sheet.title for sheet in sheets if sheet.title not in uncompared]
    _write_changes_sheet(workbook.create_sheet(CHANGES_TITLE), changes, compared, uncompared)

    saved = io.BytesIO()
    workbook.save(saved)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_ignoring_numbers_stored_as_text(saved.getvalue()))


def _write_data_sheet(sheet: Worksheet, data: Sheet) -> None:
    header = list(data.header)
    currency_index = header.index("currency") if "currency" in header else None
    currencies = [values[currency_index] if currency_index is not None else None for values in data.rows]
    _write_table(sheet, header, data.rows, currencies, top=1)
    key_index = header.index(data.key)
    for row_number, values in enumerate(data.rows, start=2):
        key = values[key_index]
        for column, name in enumerate(header, start=1):
            if key in data.added:
                sheet.cell(row=row_number, column=column).fill = ADDED_FILL
            elif (key, name) in data.changed:
                sheet.cell(row=row_number, column=column).fill = CHANGED_FILL


def _write_changes_sheet(
    sheet: Worksheet, changes: Sequence[Change], compared: Sequence[str], uncompared: Sequence[str]
) -> None:
    # The legend has the top row to itself, above the table: always in view under
    # the frozen header, outside the filtered range, and one line long, running on
    # across the empty cells beside it.
    legend = sheet.cell(row=1, column=1, value=LEGEND)
    legend.font = _NOTE_FONT
    top = 2
    rows = [[c.source, c.key, c.change, c.field, c.before, c.after] for c in changes]
    _write_table(sheet, CHANGES_HEADER, rows, [c.currency for c in changes], top=top)
    kind_column = CHANGES_HEADER.index("change") + 1
    for row_number, change in enumerate(changes, start=top + 1):
        sheet.cell(row=row_number, column=kind_column).fill = KIND_FILLS[change.change]

    # Zero changes are stated, not left as an empty table — the change report does
    # the same, because silence reads as "not checked". A first snapshot is named.
    notes = []
    if compared and not changes:
        notes.append(NO_CHANGES if not uncompared else f"No changes in {_names(compared)} since the previous snapshot.")
    if uncompared:
        notes.append(
            f"First snapshot of {_names(uncompared)}: there is no earlier one to compare with, "
            f"so nothing there is marked as new or changed."
        )
    # Under the table, and a row apart from it when it has rows, so the note is
    # never read as one of them.
    first_note = top + 1 + len(changes) + (1 if changes else 0)
    for offset, note in enumerate(notes):
        sheet.cell(row=first_note + offset, column=1, value=note).font = _NOTE_FONT


def _names(titles: Sequence[str]) -> str:
    return titles[0] if len(titles) == 1 else f"{', '.join(titles[:-1])} and {titles[-1]}"


def _write_table(
    sheet: Worksheet,
    header: Sequence[str],
    rows: Sequence[Sequence[Any]],
    currencies: Sequence[Any],
    *,
    top: int,
) -> None:
    """A header in bold on row `top`, then a row per entry: typed, formatted, sized, frozen and filtered."""
    shown = [[_shown(value, currency) for value in values] for values, currency in zip(rows, currencies, strict=True)]
    widths = []
    for column, name in enumerate(header):
        texts = [line[column] for line, values in zip(shown, rows, strict=True) if _is_text(values[column])]
        prose = any(any(character.isspace() for character in text) for text in texts)
        longest = max([len(name) + _FILTER_BUTTON, *(_longest_line(line[column]) for line in shown)])
        widths.append(min(MAX_WIDTH if prose else TOKEN_WIDTH, longest + _PADDING))

    for column, name in enumerate(header, start=1):
        sheet.cell(row=top, column=column, value=name).font = _HEADER_FONT
    for row_number, (values, texts, currency) in enumerate(zip(rows, shown, currencies, strict=True), start=top + 1):
        for column, (value, text) in enumerate(zip(values, texts, strict=True), start=1):
            cell = sheet.cell(row=row_number, column=column)
            _put(cell, value, currency, text, width=widths[column - 1], link=header[column - 1] == URL_COLUMN)

    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{top + 1}"
    sheet.auto_filter.ref = f"A{top}:{get_column_letter(len(header))}{top + len(rows)}"


def _put(cell: Cell, value: Any, currency: Any, shown: str, *, width: float, link: bool) -> None:
    """Write one value as its own type, with the format a reader expects of it."""
    cell.alignment = _TOP
    if value is None:
        return
    if isinstance(value, bool):
        cell.value = value
    elif isinstance(value, Decimal):
        cell.value = value
        cell.number_format = _money_format(currency)
    elif isinstance(value, int):
        cell.value = value
        cell.number_format = INTEGER_FORMAT
    elif isinstance(value, float):
        cell.value = value
    elif isinstance(value, datetime):
        cell.value = _naive_utc(value)
        cell.number_format = DATETIME_FORMAT
    else:
        cell.value = shown
        # Text is text, whatever it starts with: left alone, a scraped title that
        # begins with `=` would be stored as a formula and run in the reader's
        # spreadsheet — the hole the CSV export closes with a leading quote.
        cell.data_type = "s"
        if len(shown) > width - _PADDING and any(character.isspace() for character in shown):
            # Prose longer than its column wraps; an address or a slug has no
            # space to break at, and wrapping it would only chop it.
            cell.alignment = _TOP_WRAPPED
        if link and _public_page(shown):
            cell.hyperlink = shown
            cell.font = _LINK_FONT


def _public_page(text: str) -> bool:
    """Whether `text` is the address of a page a reader could open from anywhere.

    The demo store's addresses are loopback ones — they only ever answered on the
    machine that scraped them — so they stay plain text rather than dead links.
    """
    parts = urlsplit(text)
    if parts.scheme not in ("http", "https") or not parts.hostname:
        return False
    if parts.hostname == "localhost" or parts.hostname.endswith(".localhost"):
        return False
    try:
        return ipaddress.ip_address(parts.hostname).is_global
    except ValueError:
        return True  # a name rather than an address: public, as far as this file can tell


def _is_text(value: Any) -> bool:
    return value is not None and not isinstance(value, bool | int | float | Decimal | datetime)


def _longest_line(text: str) -> int:
    return max(map(len, text.splitlines() or [""]))


def _shown(value: Any, currency: Any = None) -> str:
    """A value as the cell shows it, for sizing the column and for writing text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        return f"{_SYMBOLS.get(currency, '')}{value:,.2f}"
    if isinstance(value, datetime):
        return _naive_utc(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dict):
        return "; ".join(f"{name}: {_shown(item, currency)}" for name, item in value.items())
    if isinstance(value, list | tuple):
        return ", ".join(_shown(item, currency) for item in value)
    # A control character is legal in JSON and illegal in the XML a workbook is
    # made of: dropped here, rather than failing the whole night's export.
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _money_format(currency: Any) -> str:
    symbol = _SYMBOLS.get(currency)
    return f'"{symbol}"{MONEY_FORMAT}' if symbol else MONEY_FORMAT


def _naive_utc(value: datetime) -> datetime:
    """A spreadsheet date has no time zone, so a time is stored as UTC and the header says so."""
    return value.astimezone(UTC).replace(tzinfo=None) if value.tzinfo else value


def _ignoring_numbers_stored_as_text(package: bytes) -> bytes:
    """The saved workbook, with every sheet told that digits stored as text are meant."""
    written = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(package)) as source, zipfile.ZipFile(written, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if _WORKSHEET_PART.fullmatch(item.filename):
                follower = _FOLLOWS_IGNORED_ERRORS.search(data)
                at = follower.start() if follower else data.rindex(b"</worksheet>")
                data = data[:at] + _IGNORED_ERRORS + data[at:]
            target.writestr(item, data)
    return written.getvalue()
